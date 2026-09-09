# -*- coding: utf-8 -*-
"""把 pptx 粗略渲染为 PNG，用于人工核对版式（非精确渲染）。"""
import os
import sys

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Emu

SCALE = 60.0  # px per inch
FONT_CANDIDATES = [
    ("/System/Library/Fonts/Hiragino Sans GB.ttc", 0),
    ("/System/Library/Fonts/PingFang.ttc", 0),
    ("/System/Library/Fonts/Supplemental/Songti.ttc", 0),
]


def load_font(size, bold=False):
    size = max(6, int(size))
    for path, idx in FONT_CANDIDATES:
        if os.path.exists(path):
            try:
                return ImageFont.truetype(path, size, index=1 if bold else idx)
            except Exception:
                try:
                    return ImageFont.truetype(path, size, index=idx)
                except Exception:
                    continue
    return ImageFont.load_default()


def emu_px(v):
    return int(Emu(v).inches * SCALE)


def rgb(color, default=(80, 88, 95)):
    try:
        if color and color.type is not None and color.rgb is not None:
            c = color.rgb
            return (c[0], c[1], c[2])
    except Exception:
        pass
    return default


def wrap(draw, text, font, max_w):
    lines, cur = [], ""
    for ch in text:
        test = cur + ch
        if draw.textlength(test, font=font) > max_w and cur:
            lines.append(cur)
            cur = ch
        else:
            cur = test
    lines.append(cur)
    return lines


def _defrpr_chain(shp, p):
    """返回继承链上的 defRPr 元素（段落级 → 形状 lstStyle）。"""
    from pptx.oxml.ns import qn as _qn
    out = []
    pPr = p._p.find(_qn("a:pPr"))
    if pPr is not None:
        el = pPr.find(_qn("a:defRPr"))
        if el is not None:
            out.append(el)
    try:
        txBody = shp.text_frame._txBody
        ls = txBody.find(_qn("a:lstStyle"))
        if ls is not None:
            lvl = ls.find(_qn("a:lvl1pPr"))
            if lvl is not None:
                el = lvl.find(_qn("a:defRPr"))
                if el is not None:
                    out.append(el)
    except Exception:
        pass
    return out


def resolve_style(shp, p, runs):
    from pptx.oxml.ns import qn as _qn
    # 取字符数最多的 run 作为整段代表，避免被 bullet 之类的短 run 带偏
    main = max(runs, key=lambda r: len(r.text or ""))
    ordered = [main] + [r for r in runs if r is not main]
    size = next((r.font.size.pt for r in ordered if r.font.size), None)
    color = None
    for r in ordered:
        c = rgb(r.font.color, None)
        if c:
            color = c
            break
    bold = next((bool(r.font.bold) for r in ordered if r.font.bold is not None), None)
    for el in _defrpr_chain(shp, p):
        if size is None and el.get("sz"):
            size = int(el.get("sz")) / 100.0
        if color is None:
            sc = el.find(_qn("a:solidFill"))
            if sc is not None:
                srgb = sc.find(_qn("a:srgbClr"))
                if srgb is not None:
                    v = srgb.get("val")
                    color = (int(v[0:2], 16), int(v[2:4], 16), int(v[4:6], 16))
        if bold is None and el.get("b"):
            bold = el.get("b") == "1"
    return size or 18, color or (80, 88, 95), bool(bold)


def draw_shape(dr, shp, img):
    try:
        x, y = emu_px(shp.left), emu_px(shp.top)
        w, h = emu_px(shp.width), emu_px(shp.height)
    except Exception:
        return
    st = shp.shape_type
    if shp.shape_type is not None and str(st).startswith("PICTURE"):
        try:
            blob = shp.image.blob
            import io
            pic = Image.open(io.BytesIO(blob)).convert("RGB").resize((max(1, w), max(1, h)))
            img.paste(pic, (x, y))
            dr.rectangle([x, y, x + w, y + h], outline=(200, 200, 200))
        except Exception:
            dr.rectangle([x, y, x + w, y + h], outline=(255, 0, 255))
        return
    if shp.has_text_frame is False and not hasattr(shp, "fill"):
        return
    fill = None
    try:
        if shp.fill.type is not None and str(shp.fill.type).startswith("SOLID"):
            fill = rgb(shp.fill.fore_color, (240, 240, 240))
    except Exception:
        pass
    outline = None
    try:
        if shp.line.fill.type is not None and str(shp.line.fill.type).startswith("SOLID"):
            outline = rgb(shp.line.color, (200, 200, 200))
    except Exception:
        pass
    if fill or outline:
        dr.rectangle([x, y, x + w, y + h], fill=fill, outline=outline)
    if not shp.has_text_frame:
        return
    tf = shp.text_frame
    ml = emu_px(tf.margin_left) if tf.margin_left is not None else 5
    mr = emu_px(tf.margin_right) if tf.margin_right is not None else 5
    mt = emu_px(tf.margin_top) if tf.margin_top is not None else 3
    avail = max(10, w - ml - mr)
    # 先算总高度用于垂直居中
    blocks = []
    for p in tf.paragraphs:
        runs = [r for r in p.runs]
        if not runs:
            blocks.append((None, 0, []))
            continue
        size, col, bold = resolve_style(shp, p, runs)
        px = size * SCALE / 72.0
        text = "".join(r.text for r in runs)
        font = load_font(px, bold)
        ls = p.line_spacing if isinstance(p.line_spacing, float) else 1.15
        sb = p.space_before.pt if p.space_before else 0
        sa = p.space_after.pt if p.space_after else 0
        lines = wrap(ImageDraw.Draw(Image.new("RGB", (2, 2))), text, font, avail) if text else [""]
        bh = sb * SCALE / 72.0 + len(lines) * px * ls + sa * SCALE / 72.0
        blocks.append(((font, col, px, ls, sb, sa, str(p.alignment)), bh, lines))
    if not any(b[0] for b in blocks):
        return
    total = sum(b[1] for b in blocks)
    anchor = str(tf.vertical_anchor)
    cy = y + mt
    if "MIDDLE" in anchor:
        cy = y + (h - total) / 2.0
    for meta, bh, lines in blocks:
        if meta is None:
            cy += 6
            continue
        font, col, px, ls, sb, sa, align = meta
        cy += sb * SCALE / 72.0
        for ln in lines:
            tw = dr.textlength(ln, font=font)
            tx = x + ml
            if "CENTER" in align:
                tx = x + ml + (avail - tw) / 2.0
            elif "RIGHT" in align:
                tx = x + ml + (avail - tw)
            dr.text((tx, cy), ln, font=font, fill=col)
            cy += px * ls
        cy += sa * SCALE / 72.0
    # 溢出提示
    if cy > y + h + 3:
        dr.rectangle([x, y, x + w, y + h], outline=(255, 0, 0), width=3)
        dr.text((x + 4, y + h - 18), "OVERFLOW", font=load_font(16), fill=(255, 0, 0))


def main(path, outdir):
    prs = Presentation(path)
    W = int(prs.slide_width.inches * SCALE)
    H = int(prs.slide_height.inches * SCALE)
    os.makedirs(outdir, exist_ok=True)
    for i, slide in enumerate(prs.slides, 1):
        img = Image.new("RGB", (W, H), (255, 255, 255))
        dr = ImageDraw.Draw(img)
        for shp in slide.slide_layout.shapes:
            draw_shape(dr, shp, img)
        for shp in slide.shapes:
            draw_shape(dr, shp, img)
        dr.rectangle([0, 0, W - 1, H - 1], outline=(0, 0, 0))
        img.save(os.path.join(outdir, "s%02d.png" % i))
    print("rendered", len(prs.slides._sldIdLst), "->", outdir)


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
